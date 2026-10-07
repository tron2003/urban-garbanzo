"""Engine behaviour, Excel<->code validation, and the 'add an 11th workflow with minimal changes' claim."""
import csv
import shutil

import pytest
from openpyxl import load_workbook

from workflow_agent.agent import Agent
from workflow_agent.core.errors import NeedMoreInput, StepError, ToolNotPermitted, WorkflowConfigError
from workflow_agent.core.models import Param, Table
from workflow_agent.core.registry import DEFAULT_REGISTRY, WorkflowModule
from workflow_agent.llm import OfflineLLM
from tests.conftest import ROOT, make_settings

NEW_ROW = ["WF011", "Product Count", "User asks how many products are in the catalogue", "Product CSV",
           "Load products → count products → report", "If the count is zero, warn", "CSV reader", "Number of products"]


@pytest.fixture
def wb_with_wf011(tmp_path):
    dst = tmp_path / "workflows.xlsx"
    shutil.copy(ROOT / "data" / "AI_Agent_Workflow_Assessment.xlsx", dst)
    wb = load_workbook(dst)
    wb["Workflows"].append(NEW_ROW)
    wb["Test_Questions"].append(["WF011", "How many products do we have?", "extensibility"])
    wb.save(dst)
    return dst


@pytest.fixture
def cleanup():
    yield
    DEFAULT_REGISTRY.unregister("WF011")


def define_wf011():
    wf = WorkflowModule("WF011", params=[Param("product_file", "Product CSV", type="file", sample_default="data/products.csv")],
                        routing_hints=["how many products", "product count", "size of the catalogue"])

    @wf.step("Load products")
    def load(ctx):
        ctx.state["table"] = ctx.call_tool("csv_reader", op="read", path=ctx.param("product_file"))
        return f"{len(ctx.state['table']['rows'])} rows"

    @wf.step("count products")
    def count(ctx):
        ctx.state["n"] = len(ctx.state["table"]["rows"])
        ctx.decide("count is zero", "no" if ctx.state["n"] else "yes")

    @wf.step("report")
    def report(ctx):
        ctx.set_result(f"There are {ctx.state['n']} products.", tables=[Table("Count", ["Products"], [[ctx.state["n"]]])],
                       data={"count": ctx.state["n"]})
    return wf


# ---- the headline extensibility claim ------------------------------------------------------
def test_adding_an_eleventh_workflow_needs_only_an_excel_row_and_one_module(tmp_path, wb_with_wf011, cleanup):
    define_wf011()                                              # the ONLY code: one new module
    s = make_settings(tmp_path, workflow_excel=wb_with_wf011)
    agent = Agent(s, llm=OfflineLLM())                          # strict startup validation passes
    assert len(agent.specs) == 11
    resp = agent.handle("How many products do we have?")       # routed from Excel text + hints, no router edits
    assert resp.report.workflow_id == "WF011" and resp.report.result.data["count"] == 12
    assert [x.status for x in resp.report.steps] == ["success"] * 3
    assert agent.handle("Which products need restocking?").report.workflow_id == "WF001"      # old ones unaffected


def test_excel_row_without_python_fails_fast_with_instructions(tmp_path, wb_with_wf011):
    with pytest.raises(WorkflowConfigError) as e:
        Agent(make_settings(tmp_path, workflow_excel=wb_with_wf011), llm=OfflineLLM())
    assert "WF011" in str(e.value) and "WorkflowModule('WF011')" in str(e.value)


def test_renamed_excel_step_is_detected(tmp_path, wb_with_wf011, cleanup):
    define_wf011()
    wb = load_workbook(wb_with_wf011); ws = wb["Workflows"]
    ws.cell(row=ws.max_row, column=5).value = "Load products → tally products → report"
    wb.save(wb_with_wf011)
    with pytest.raises(WorkflowConfigError) as e:
        Agent(make_settings(tmp_path, workflow_excel=wb_with_wf011), llm=OfflineLLM())
    msg = str(e.value)
    assert "'tally products' has no @step handler" in msg and "does not match any Excel step" in msg


def test_unknown_excel_tool_is_detected(tmp_path, wb_with_wf011, cleanup):
    define_wf011()
    wb = load_workbook(wb_with_wf011); ws = wb["Workflows"]
    ws.cell(row=ws.max_row, column=7).value = "CSV reader; quantum oracle"
    wb.save(wb_with_wf011)
    with pytest.raises(WorkflowConfigError, match="quantum oracle"):
        Agent(make_settings(tmp_path, workflow_excel=wb_with_wf011), llm=OfflineLLM())


def test_orphan_python_module_is_only_a_warning(tmp_path, cleanup):
    define_wf011()
    agent = Agent(make_settings(tmp_path), llm=OfflineLLM())
    assert any("WF011" in w for w in agent.config_warnings)


# ---- engine robustness ---------------------------------------------------------------------
def test_tool_not_listed_in_excel_is_blocked(agent):
    # WF001's Excel row lists CSV reader + calculator only
    ctx_spec = agent.spec_by_id["WF001"]
    from workflow_agent.core.context import RunContext
    ctx = RunContext(ctx_spec, {}, {}, agent.settings, agent.tools, agent.llm, "t")
    with pytest.raises(ToolNotPermitted, match="not allowed"):
        ctx.call_tool("order_api", op="search", order_id="ORD-1001")
    assert ctx.tool_calls[-1].status == "error"


def test_unexpected_exception_does_not_crash_the_agent_and_skips_remaining_steps(agent, monkeypatch):
    from workflow_agent.workflows import wf001_inventory_restock as m
    monkeypatch.setitem(m.wf.handlers, "compare current stock with minimum threshold", lambda ctx: 1 / 0)
    r = agent.run_workflow("WF001")
    assert r.status == "failed" and "ZeroDivisionError" in r.steps[1].error
    assert [s.status for s in r.steps] == ["success", "failed", "skipped", "skipped", "skipped"]
    assert r.result is None


def test_workflow_that_forgets_to_set_a_result_is_a_failure(agent, monkeypatch):
    from workflow_agent.workflows import wf001_inventory_restock as m
    monkeypatch.setitem(m.wf.handlers, "generate restock list", lambda ctx: "did nothing")
    r = agent.run_workflow("WF001")
    assert r.status == "failed" and "no result" in r.error


def test_run_log_is_written_in_the_wf010_schema(agent, settings):
    agent.run_workflow("WF001"); agent.run_workflow("WF005", {"order_id": "ORD-9999"})
    rows = list(csv.DictReader(settings.log_file.open(encoding="utf-8")))
    assert {r["workflow_id"] for r in rows} == {"WF001", "WF005"}
    wf5 = [r for r in rows if r["workflow_id"] == "WF005"]
    assert [r["step_status"] for r in wf5] == ["success", "needs_input", "skipped", "skipped", "skipped"]
    assert all(r["run_status"] == "needs_input" for r in wf5)
    assert set(rows[0]) == {"timestamp", "run_id", "workflow_id", "workflow_name", "step_name", "step_status",
                            "duration_ms", "run_status", "error_type", "error_message"}


def test_unwritable_log_never_breaks_a_run(agent, settings, tmp_path):
    blocker = tmp_path / "file"; blocker.write_text("x")
    agent.engine.logger.path = blocker / "sub" / "log.csv"       # parent is a file -> OSError
    assert agent.run_workflow("WF001").status == "success"


def test_every_workflow_runs_end_to_end_without_arguments(agent):
    """Eight succeed with bundled sample inputs; WF005 must ask for an order id and WF007 for goal/dates (Excel rules)."""
    statuses = {w: agent.run_workflow(w).status for w in agent.spec_by_id}
    assert statuses.pop("WF007") == "needs_input" and statuses.pop("WF005") == "needs_input"
    assert set(statuses.values()) == {"success"} and len(statuses) == 8
