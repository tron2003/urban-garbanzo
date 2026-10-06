"""Reads workflow definitions from the Excel workbook.

The workbook is the *source of truth* for: workflow id/name, trigger, inputs,
the ordered step list, decision logic, required tools and expected output.
Python only supplies the executable binding for each step (see registry.py).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .core.errors import WorkflowConfigError

REQUIRED_COLUMNS = [
    "workflow_id", "workflow_name", "trigger", "inputs",
    "steps", "decision_logic", "tools_required", "expected_output",
]
_STEP_SPLIT = re.compile(r"\s*(?:→|->|=>|›|»)\s*")


def norm_key(text: str) -> str:
    """Canonical form used to match Excel step text to Python step handlers."""
    return re.sub(r"\s+", " ", str(text or "").strip().lower())


def _col(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(header or "").strip().lower()).strip("_")


@dataclass(frozen=True)
class WorkflowSpec:
    id: str
    name: str
    trigger: str
    inputs: str
    steps: tuple[str, ...]
    decision_logic: str
    tools: tuple[str, ...]
    expected_output: str
    examples: tuple[str, ...] = field(default_factory=tuple)   # from the Test_Questions sheet


def load_workflow_specs(path: str | Path) -> list[WorkflowSpec]:
    path = Path(path)
    if not path.exists():
        raise WorkflowConfigError(f"Workflow workbook not found: {path}")

    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # corrupt / not an xlsx
        raise WorkflowConfigError(f"Cannot open workbook {path.name}: {exc}") from exc

    sheet = next((wb[n] for n in wb.sheetnames if _col(n) == "workflows"), None)
    if sheet is None:
        raise WorkflowConfigError(
            f"Sheet 'Workflows' not found in {path.name} (sheets: {', '.join(wb.sheetnames)})")

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise WorkflowConfigError("Sheet 'Workflows' is empty")

    headers = [_col(h) for h in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise WorkflowConfigError(f"Sheet 'Workflows' is missing column(s): {', '.join(missing)}")
    idx = {h: i for i, h in enumerate(headers) if h}

    examples: dict[str, list[str]] = {}
    tq = next((wb[n] for n in wb.sheetnames if _col(n) == "test_questions"), None)
    if tq is not None:
        tq_rows = list(tq.iter_rows(values_only=True))
        if tq_rows:
            th = [_col(h) for h in tq_rows[0]]
            if "workflow_id" in th and "test_request" in th:
                wi, ri = th.index("workflow_id"), th.index("test_request")
                for r in tq_rows[1:]:
                    wid = str(r[wi] or "").strip()
                    req = str(r[ri] or "").strip()
                    if wid and req:
                        examples.setdefault(wid, []).append(req)

    problems: list[str] = []
    specs: list[WorkflowSpec] = []
    seen: set[str] = set()
    for line_no, r in enumerate(rows[1:], start=2):
        get = lambda name: str(r[idx[name]]).strip() if r[idx[name]] is not None else ""
        wid = get("workflow_id")
        if not wid:
            if any(c not in (None, "") for c in r):
                problems.append(f"Row {line_no}: Workflow_ID is empty")
            continue
        if wid in seen:
            problems.append(f"Row {line_no}: duplicate Workflow_ID {wid}")
            continue
        seen.add(wid)
        steps = tuple(s.strip() for s in _STEP_SPLIT.split(get("steps")) if s.strip())
        if not steps:
            problems.append(f"{wid}: Steps column is empty")
        if not get("workflow_name"):
            problems.append(f"{wid}: Workflow_Name is empty")
        tools = tuple(t.strip() for t in re.split(r"[;\n]", get("tools_required")) if t.strip())
        specs.append(WorkflowSpec(
            id=wid, name=get("workflow_name"), trigger=get("trigger"), inputs=get("inputs"),
            steps=steps, decision_logic=get("decision_logic"), tools=tools,
            expected_output=get("expected_output"), examples=tuple(examples.get(wid, ())),
        ))

    if problems:
        raise WorkflowConfigError(problems)
    if not specs:
        raise WorkflowConfigError("No workflows found in the workbook")
    return specs
