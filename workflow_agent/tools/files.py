"""File readers: CSV, Excel and (read-only) SQLite."""
from __future__ import annotations

import csv
import io
import re
import sqlite3
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from ..core.errors import ToolError
from ..utils import is_blank, map_columns, snake
from .registry import tool


def _resolve(svc, path: str) -> Path:
    p = svc.settings.resolve(path)
    if not p.exists():
        raise ToolError("file", f"File not found: {path}")
    if p.is_dir():
        raise ToolError("file", f"Expected a file but got a folder: {path}")
    return p


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_csv_text(text: str) -> list[list[str]]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _rows_to_records(matrix: list[list[str]], header_idx: int, source: str) -> dict:
    headers = [h.strip() for h in matrix[header_idx]]
    # de-duplicate / fill blank header names so dict keys stay unique
    seen: dict[str, int] = {}
    clean = []
    for i, h in enumerate(headers):
        name = h or f"column_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        clean.append(name)
    rows, blank = [], 0
    for offset, raw in enumerate(matrix[header_idx + 1:], start=header_idx + 2):
        if all(is_blank(c) for c in raw):
            blank += 1
            continue
        rec = {clean[i]: (raw[i] if i < len(raw) else "") for i in range(len(clean))}
        rec["_row"] = offset
        rows.append(rec)
    return {"headers": clean, "rows": rows, "blank_rows": blank, "source": source,
            "header_row": header_idx + 1, "preamble_rows": header_idx}


@tool("csv_reader", aliases=("CSV reader",), description="Read a CSV/TSV file into records.")
def csv_reader(svc, op: str = "read", path: str = "") -> dict:
    if op != "read":
        raise ToolError("csv_reader", f"unsupported op '{op}'")
    p = _resolve(svc, path)
    if p.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        raise ToolError("csv_reader", f"{p.name} is an Excel file; this workflow expects CSV")
    matrix = read_csv_text(_decode(p.read_bytes()))
    matrix = [r for r in matrix if r is not None]
    if not matrix or all(is_blank(c) for c in matrix[0]):
        raise ToolError("csv_reader", f"{p.name} is empty")
    return _rows_to_records(matrix, 0, p.name)


@tool("csv_database_reader", aliases=("CSV/database reader",),
      description="Read a CSV file or a SQLite table (read-only).")
def csv_database_reader(svc, op: str = "read", path: str = "", table: Optional[str] = None,
                        sql: Optional[str] = None) -> dict:
    p = _resolve(svc, path)
    if p.suffix.lower() in (".db", ".sqlite", ".sqlite3"):
        if not sql and not table:
            raise ToolError("csv_database_reader", "SQLite sources need 'table' or 'sql'")
        query = sql or f'SELECT * FROM "{re.sub(r"[^A-Za-z0-9_]", "", table)}"'
        if not re.match(r"^\s*select\b", query, re.I):
            raise ToolError("csv_database_reader", "only SELECT statements are allowed")
        try:
            con = sqlite3.connect(f"file:{p}?mode=ro", uri=True)
            con.row_factory = sqlite3.Row
            rows = [dict(r) for r in con.execute(query)]
            con.close()
        except sqlite3.Error as exc:
            raise ToolError("csv_database_reader", f"database error: {exc}") from exc
        for i, r in enumerate(rows, start=2):
            r["_row"] = i
        return {"headers": list(rows[0].keys()) if rows else [], "rows": rows, "blank_rows": 0,
                "source": p.name, "header_row": 1, "preamble_rows": 0}
    return csv_reader(svc, op="read", path=path)


@tool("excel_csv_parser", aliases=("Excel/CSV parser",),
      description="Parse CSV or XLSX vendor files; detects the header row when there is a preamble.")
def excel_csv_parser(svc, op: str = "read", path: str = "", sheet: Optional[str] = None,
                     header_hints: Optional[dict] = None) -> dict:
    p = _resolve(svc, path)
    ext = p.suffix.lower()
    if ext in (".csv", ".tsv", ".txt"):
        matrix = read_csv_text(_decode(p.read_bytes()))
        sheet_name = ""
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(p, read_only=True, data_only=True)
        except Exception as exc:
            raise ToolError("excel_csv_parser", f"cannot open {p.name}: {exc}") from exc
        if sheet and sheet not in wb.sheetnames:
            raise ToolError("excel_csv_parser", f"sheet '{sheet}' not found (available: {', '.join(wb.sheetnames)})")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        sheet_name = ws.title
        matrix = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
    elif ext == ".xls":
        raise ToolError("excel_csv_parser", "legacy .xls is not supported; save the file as .xlsx or .csv")
    else:
        raise ToolError("excel_csv_parser", f"unsupported file type '{ext or 'unknown'}' (use .csv or .xlsx)")

    # drop trailing empty rows but keep leading ones for header detection
    while matrix and all(is_blank(c) for c in matrix[-1]):
        matrix.pop()
    if not matrix:
        raise ToolError("excel_csv_parser", f"{p.name} contains no data")

    header_idx = 0
    if header_hints:
        best, best_hits = 0, 0
        for i, row in enumerate(matrix[:15]):
            hits = len(map_columns([c for c in row if c], header_hints))
            if hits > best_hits:
                best, best_hits = i, hits
        if best_hits >= 2:
            header_idx = best
    out = _rows_to_records(matrix, header_idx, p.name)
    out["sheet"] = sheet_name
    return out
